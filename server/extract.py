"""正文提取 worker：pending 文件 → 提取文本 → zlib 压缩入 contents + 更新 FTS。

UI 永不直接展示正文；正文只作检索/嵌入燃料。
扫描件（文本极少且图片多）打 scanned 标记，MVP 不做 OCR。
"""
import zlib

from . import config, ingest, tags
from .schema import connect

MAX_TEXT_CHARS = 1_000_000
SCANNED_TEXT_MIN = 100


def _read_text_file(path: str) -> tuple[str, int]:
    for enc in ("utf-8", "gb18030"):
        try:
            with open(path, encoding=enc) as f:
                return f.read(MAX_TEXT_CHARS), 0
        except UnicodeDecodeError:
            continue
    # 兜底：坏字节替换，不让文件整体提取失败（大概率是 utf-16 或二进制混杂）
    with open(path, encoding="utf-8", errors="replace") as f:
        return f.read(MAX_TEXT_CHARS), 0


def _extract_pdf(path: str) -> tuple[str, int]:
    import fitz
    parts, images = [], 0
    with fitz.open(path) as doc:
        for page in doc:
            parts.append(page.get_text())
            images += len(page.get_images(full=True))
            if sum(len(p) for p in parts) > MAX_TEXT_CHARS:
                break
    return "\n".join(parts), images


def _ooxml_text_fallback(path: str, member_prefix: str) -> tuple[str, int]:
    """python-docx/pptx 打不开的坏包（WPS 常见）兜底：直接从 zip 抠文本节点。"""
    import re
    import zipfile
    parts = []
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if not n.startswith(member_prefix) or not n.endswith(".xml"):
                continue
            xml = z.read(n).decode("utf-8", "replace")
            parts += re.findall(r"<(?:w|a):t[^>]*>([^<]*)</(?:w|a):t>", xml)
    return "".join(parts), 0


def _extract_docx(path: str) -> tuple[str, int]:
    import docx
    try:
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for tbl in d.tables:
            for row in tbl.rows:
                parts.append(" ".join(c.text for c in row.cells))
        return "\n".join(parts), len(d.inline_shapes)
    except Exception:
        return _ooxml_text_fallback(path, "word/")


def _extract_pptx(path: str) -> tuple[str, int]:
    from pptx import Presentation
    try:
        prs = Presentation(path)
        parts, images = [], 0
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
                if shape.shape_type == 13:  # PICTURE
                    images += 1
            if slide.has_notes_slide:
                ntf = slide.notes_slide.notes_text_frame
                if ntf is not None:
                    parts.append(ntf.text)
        return "\n".join(parts), images
    except Exception:
        return _ooxml_text_fallback(path, "ppt/slides/")


def _extract_xlsx(path: str) -> tuple[str, int]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(ws.title)
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if isinstance(cell, str) and cell.strip():
                    parts.append(cell)
    wb.close()
    return "\n".join(parts), 0


EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".pptx": _extract_pptx,
    ".xlsx": _extract_xlsx,
    ".md": _read_text_file,
    ".txt": _read_text_file,
}


def extract_one(conn, vol: str, frn: int) -> str:
    """提取单个文件。返回新 status。"""
    row = conn.execute(
        "SELECT path, ext FROM files WHERE volume=? AND frn=?",
        (vol, frn)).fetchone()
    if not row:
        return "gone"
    path, ext = row

    st_mode_attrs = _stat_attrs(path)
    if st_mode_attrs is None:
        status = "failed"          # 文件已不存在/无权限
    elif st_mode_attrs & (config.FILE_ATTRIBUTE_RECALL_ON_DATA_ACCESS |
                          config.FILE_ATTRIBUTE_OFFLINE):
        status = "cloud"           # OneDrive 占位，不触发下载
    elif ext in EXTRACTORS:
        try:
            text, images = EXTRACTORS[ext](path)
            text = text[:MAX_TEXT_CHARS]  # 统一截断，台账类 xlsx 可能解出 GB 级
            scanned = 1 if (ext == ".pdf" and images > 0 and
                            len(text.strip()) < SCANNED_TEXT_MIN) else 0
            conn.execute(
                "INSERT OR REPLACE INTO contents(file_pk,text_zst) "
                "VALUES(?,?)", (config.pk(vol, frn), zlib.compress(
                    text.encode("utf-8"), 6)))
            ingest.fts_delete(conn, config.pk(vol, frn))
            ingest.fts_insert(conn, config.pk(vol, frn), path, text)
            _pk = config.pk(vol, frn)
            conn.execute("DELETE FROM tags WHERE file_pk=?", (_pk,))
            conn.executemany(
                "INSERT OR IGNORE INTO tags(file_pk, tag) VALUES(?,?)",
                [(_pk, t) for t in tags.extract_tags(text)])
            status = "extracted"
        except Exception:
            status = "failed"
        # image_count / scanned 更新放在 status 确定后
    else:
        status = "failed"          # .doc/.xls/.ppt 等旧格式，MVP 不支持

    images_now = locals().get("images", 0) if status == "extracted" else 0
    conn.execute(
        "UPDATE files SET status=?, image_count=?, scanned=CASE WHEN ?=1 "
        "THEN 1 ELSE scanned END WHERE volume=? AND frn=?",
        (status, images_now, scanned if status == "extracted" else 0,
         vol, frn))
    conn.commit()
    return status


def _stat_attrs(path: str):
    import os
    try:
        return os.stat(path).st_file_attributes
    except OSError:
        return None


def run_batch(conn, limit: int = 200) -> int:
    """处理一批 pending 文件，返回处理数。"""
    rows = conn.execute(
        "SELECT volume, frn FROM files WHERE status='pending' LIMIT ?",
        (limit,)).fetchall()
    for vol, frn in rows:
        extract_one(conn, vol, frn)
    return len(rows)


if __name__ == "__main__":
    conn = connect()
    n = 0
    while True:
        done = run_batch(conn)
        n += done
        if not done:
            break
        print(f"extracted {n} ...", flush=True)
    print(f"done: {n} files processed")
